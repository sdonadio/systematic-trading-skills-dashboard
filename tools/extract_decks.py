#!/usr/bin/env python3
"""Extract slide text from the FINM 33500 decks into deck_<key>.txt files."""
import os, sys
from pptx import Presentation

DECK_DIR = "/Users/sdonadio/PycharmProjects/AlgoArena/course/systematic-trading"
OUT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECKS = {"w0": "week0_smoketest.pptx"}
for i in range(1, 10):
    DECKS["w%d" % i] = "week%d.pptx" % i

def text_of(shape):
    out = []
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                out.append(("  " * p.level) + t)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            out.append(" | ".join(cells))
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            out.extend(text_of(s))
    return out

for key, fn in DECKS.items():
    path = os.path.join(DECK_DIR, fn)
    prs = Presentation(path)
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append("--- %s slide %d ---" % (key, i))
        for sh in slide.shapes:
            lines.extend(text_of(sh))
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame.text.strip()
            if nt:
                lines.append("[notes] " + nt.replace("\n", " / "))
    open(os.path.join(OUT, "deck_%s.txt" % key), "w").write("\n".join(lines) + "\n")
    print("%-4s %-24s %3d slides" % (key, fn, len(prs.slides.__iter__().__length_hint__() if False else list(prs.slides))))
